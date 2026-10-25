"""Fixes from the review of everything built so far: numbers too big for the database, removing
people who leave, the daily run and renewals surviving a failure, spreadsheets that can't run
formulas, completed forms staying complete, live location only while clocked in, and dashboards
that don't load every job the company ever had."""
import io
from datetime import timedelta

from openpyxl import load_workbook

from app import db
from app.models import CustomerMessage, Job, JobEvent, Payment, Review, ServiceRequest, Subscription, User
from app.services.billing import Event, TestProvider, mechanic_count, renew_due
from app.services.data_export import cell
from app.utils.time import utcnow
from conftest import login


def boss(auth):
    return auth("boss@acme.test")


# ---------- Numbers too big for the database ----------


def test_numbers_too_big_for_the_database_are_refused_not_crashed(client, auth, data):
    """They used to reach the database and fail there (a 500): an ID in the address is simply not
    found, and one in the request is a clear error."""
    big = 10**20
    for path in (f"/api/jobs/{big}", f"/api/machines/{big}", f"/api/files/{big}", f"/api/customers/{big}"):
        assert client.get(path, headers=boss(auth)).status_code == 404, path
    for path, body in (
        ("/api/jobs", {"title": "Fix", "description": "Broken", "machine_id": big}),
        ("/api/sites", {"customer_id": big, "name": "Yard"}),
        ("/api/catalog", {"name": "Service", "checklist_id": big}),
        ("/api/routes", {"mechanic_id": big}),
    ):
        response = client.post(path, headers=boss(auth), json=body)
        assert response.status_code == 400 and "too large" in response.get_json()["error"], (path, response.get_json())


# ---------- Removing people who leave ----------


def test_removing_a_technician_with_history_ends_their_access_and_keeps_the_history(client, auth, data):
    # Mech One still has an open job: that comes first.
    refused = client.delete(f"/api/users/{data['mechanic']}", headers=boss(auth))
    assert refused.status_code == 409 and "open jobs" in refused.get_json()["error"]
    job = db.session.get(Job, data["job"])
    job.status = "Completed"
    db.session.commit()
    before = mechanic_count(data["acme"])
    signed_in = auth("mech@acme.test")

    removed = client.delete(f"/api/users/{data['mechanic']}", headers=boss(auth))
    assert removed.status_code == 200 and removed.get_json() == {"ok": True, "kept": True}
    # Signed out everywhere, and no way back in: the password and the address are gone.
    assert client.get("/api/users/me", headers=signed_in).status_code == 401
    assert login(client, "mech@acme.test").status_code == 401
    assert client.post("/api/auth/forgot-password", json={"email": "mech@acme.test"}).status_code in (200, 202)
    person = db.session.get(User, data["mechanic"])
    assert person.removed_at and person.email != "mech@acme.test" and person.name == "Mech One"
    # Left out of the lists and of the plan's technician count; can't be given work.
    listed = [row["name"] for row in client.get("/api/users/mechanics", headers=boss(auth)).get_json()["mechanics"]]
    assert "Mech One" not in listed and "Mech Two" in listed
    assert mechanic_count(data["acme"]) == before - 1
    new_job = client.post("/api/jobs", headers=boss(auth), json={"title": "Next", "description": "d", "mechanic_id": data["mechanic"]})
    assert new_job.status_code == 404
    # The history still names them, and the address can be invited again.
    detail = client.get(f"/api/jobs/{data['job']}", headers=boss(auth)).get_json()["job"]
    assert detail["mechanic"]["name"] == "Mech One"
    invited = client.post("/api/customers/invite", headers=boss(auth), json={"name": "Mech One", "email": "mech@acme.test", "role": "mechanic"})
    assert invited.status_code == 201, invited.get_json()
    # Nothing else can be done to them either: no reset code, no edits.
    assert client.delete(f"/api/users/{data['mechanic']}", headers=boss(auth)).status_code == 404
    assert client.post(f"/api/users/{data['mechanic']}/reset-code", headers=boss(auth)).status_code == 404
    assert client.post(f"/api/users/{data['mechanic']}/reset-two-factor", headers=boss(auth)).status_code == 404
    assert client.patch(f"/api/users/{data['mechanic']}", headers=boss(auth), json={"name": "Back"}).status_code == 404


def test_someone_without_history_is_deleted(client, auth, data):
    removed = client.delete(f"/api/users/{data['mechanic2']}", headers=boss(auth))
    assert removed.status_code == 200 and removed.get_json() == {"ok": True, "kept": False}
    assert db.session.get(User, data["mechanic2"]) is None


def test_a_customer_login_that_wrote_in_the_chat_can_be_removed(client, auth, data):
    """Used to fail on the message's link to them (a 500): now the login ends, the message stays."""
    db.session.add(CustomerMessage(company_id=data["acme"], customer_id=data["acme_customer"], sender_id=data["customer"], from_customer=True, body="Hello"))
    db.session.commit()
    removed = client.delete(f"/api/users/{data['customer']}", headers=boss(auth))
    assert removed.status_code == 200 and removed.get_json()["kept"] is True
    assert CustomerMessage.query.filter_by(sender_id=data["customer"]).count() == 1
    assert login(client, "cust@acme.test").status_code == 401


# ---------- The daily run and renewals survive a failure ----------


def test_the_daily_run_finishes_every_part_when_one_fails(app, client, data, monkeypatch):
    import app.routes.maintenance as maintenance

    def down():
        raise RuntimeError("payment service down")

    monkeypatch.setattr(maintenance, "renew_due", down)
    removed = maintenance.run_daily()
    assert removed["failed"] == ["billing"] and "deleted_companies" in removed and "billing" not in removed
    app.config["CRON_SECRET"] = "s3cret-value"
    response = client.get("/api/cron/daily", headers={"Authorization": "Bearer s3cret-value"})
    assert response.status_code == 500 and response.get_json()["removed"]["failed"] == ["billing"]


class UndecidedProvider(TestProvider):
    """A renewal the provider doesn't decide straight away (an M-Pesa prompt)."""

    def __init__(self):
        self.charges = 0
        self.answer = None

    def charge_saved(self, payment, subscription, email):
        self.charges += 1
        return None

    def confirm(self, payment):
        return Event(self.answer, payment.reference) if self.answer else None


def due_for_renewal(company_id):
    subscription = Subscription(
        company_id=company_id, plan="professional", interval="month", status="active",
        current_period_end=utcnow() + timedelta(hours=12), payment_token="test_card_4242", payment_method="Test card •••• 4242",
    )
    db.session.add(subscription)
    db.session.commit()
    return subscription


def test_a_renewal_still_undecided_is_asked_about_not_charged_again(app, data):
    subscription = due_for_renewal(data["acme"])
    end = subscription.current_period_end
    checkout = UndecidedProvider()
    app.extensions["billing"] = checkout
    assert renew_due() == {"charged": 0, "failed": 0, "reminded": 0, "errors": 0}
    assert checkout.charges == 1 and Payment.query.filter_by(kind="renewal", status="pending").count() == 1
    # The next day the provider still doesn't know: no second charge.
    renew_due(utcnow() + timedelta(days=1))
    assert checkout.charges == 1
    # Then it says it was paid: the plan runs on, and still nothing is charged twice.
    checkout.answer = "paid"
    assert renew_due(utcnow() + timedelta(days=2))["charged"] == 1
    db.session.refresh(subscription)
    assert subscription.current_period_end > end and checkout.charges == 1


def test_a_payment_service_error_doesnt_stop_the_other_renewals(app, data):
    due_for_renewal(data["acme"])

    class Down(TestProvider):
        def charge_saved(self, payment, subscription, email):
            raise ConnectionError("no answer")

    app.extensions["billing"] = Down()
    assert renew_due()["errors"] == 1
    # Nothing half-made is left behind: tomorrow's run tries again.
    assert Payment.query.filter_by(kind="renewal").count() == 0


# ---------- Spreadsheets can't run formulas typed into the app ----------


def test_text_that_starts_a_formula_stays_text_in_spreadsheets(client, auth, data):
    assert cell("=HYPERLINK(\"http://evil\",\"Click\")") == "'=HYPERLINK(\"http://evil\",\"Click\")"
    assert [cell(value) for value in ("+254 700", "-5", "@sum", "Fine", 42, None)] == ["'+254 700", "'-5", "'@sum", "Fine", 42, None]
    job = db.session.get(Job, data["job"])
    job.title = '=HYPERLINK("http://evil","Click")'
    job.status = "Completed"
    db.session.add(JobEvent(job_id=job.id, user_id=data["mechanic"], action="status", from_status="In Progress", to_status="Completed"))
    db.session.commit()
    response = client.get("/api/insights/export.xlsx", headers=boss(auth))
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.data))
    found = [c for sheet in book.worksheets for row in sheet.iter_rows() for c in row if c.value == job.title]
    assert found and all(c.data_type == "s" for c in found)


# ---------- Live location only while clocked in ----------


def test_live_location_needs_an_open_shift(client, auth, data):
    job = Job(title="Pump", description="d", company_id=data["acme"], customer_id=data["acme_customer"], mechanic_id=data["mechanic2"], status="In Progress")
    db.session.add(job)
    db.session.commit()
    response = client.post("/api/location/live", headers=auth("mech2@acme.test"), json={"job_id": job.id, "latitude": -1.28, "longitude": 36.82, "accuracy": 10})
    assert response.status_code == 409 and "Clock in" in response.get_json()["error"]


# ---------- Messages in the person's language ----------


def test_the_confirm_your_email_message_is_translated(client, auth, data):
    manager = db.session.get(User, data["boss"])
    manager.email_verified_at = None
    db.session.commit()
    headers = {**boss(auth), "X-Language": "sw"}
    response = client.post("/api/customers/invite", headers=headers, json={"name": "New", "email": "new@acme.test", "role": "mechanic"})
    assert response.status_code == 403 and response.get_json()["error"].startswith("Thibitisha anwani yako")


# ---------- Dashboards don't load every job ever ----------


def test_the_dashboard_lists_and_counts_come_from_the_server(client, auth, data):
    """The manager's dashboard asks every 30 seconds: it gets the unfinished jobs and the counts,
    not the whole history, and a declined request isn't counted as open."""
    for status, mechanic in (("Completed", "mechanic"), ("Cancelled", "mechanic"), ("Paused", "mechanic2"), ("Pending", None)):
        db.session.add(Job(title=status, description="d", company_id=data["acme"], customer_id=data["acme_customer"], mechanic_id=data[mechanic] if mechanic else None, status=status))
    for status in ("Submitted", "Declined", "Completed"):
        db.session.add(ServiceRequest(company_id=data["acme"], customer_id=data["acme_customer"], machine_id=data["machine"], subject=status, description="d", status=status))
    db.session.add(Review(company_id=data["acme"], customer_id=data["acme_customer"], mechanic_id=data["mechanic"], machine_id=data["machine"], service_record_id=data["record"], rating=4))
    db.session.commit()

    open_jobs = client.get("/api/jobs?open=1", headers=boss(auth)).get_json()["jobs"]
    assert sorted(job["status"] for job in open_jobs) == ["Assigned", "Paused", "Pending"]
    mine = client.get(f"/api/jobs?mechanic_id={data['mechanic']}", headers=boss(auth)).get_json()["jobs"]
    assert sorted(job["title"] for job in mine) == ["Cancelled", "Completed", "Fix generator"]
    assert client.get(f"/api/reviews?mechanic_id={data['mechanic2']}", headers=boss(auth)).get_json()["reviews"] == []
    assert len(client.get(f"/api/reviews?mechanic_id={data['mechanic']}", headers=boss(auth)).get_json()["reviews"]) == 1
    # A technician only ever sees their own, whatever they ask for.
    asked = client.get(f"/api/jobs?mechanic_id={data['mechanic2']}", headers=auth("mech@acme.test")).get_json()["jobs"]
    assert {job["mechanic_id"] for job in asked} == {data["mechanic"]}
    requests = client.get("/api/customers/requests?open=1", headers=boss(auth)).get_json()["requests"]
    assert [row["subject"] for row in requests] == ["Submitted"]

    report = client.get("/api/reports", headers=boss(auth)).get_json()
    assert report["jobs"] == {"total": 5, "completed": 1, "pending": 2, "in_progress": 1, "cancelled": 1}
    assert report["totals"]["open_requests"] == 1 and report["rating"] == 4.0 and report["review_count"] == 1
    assert client.get("/api/reports", headers=auth("cust@acme.test")).get_json()["totals"]["open_requests"] == 1
    assert client.get("/api/jobs?mechanic_id=abc", headers=boss(auth)).status_code == 400
