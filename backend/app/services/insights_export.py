"""Insights as an Excel workbook (a sheet per table, for working with the numbers) and as a PDF
(the same page as the screen, for sending on). Both are written in the reader's language, with
times on the company's clock."""
import io
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.customer import Customer
from app.models.machine import Machine
from app.models.user import User
from app.services.report_pdf import ACCENT, INK, LINE, MUTED, clean
from app.i18n import MONTHS, N_, _

HEAD_FILL = PatternFill("solid", fgColor="EAF3FF")
GENERAL = N_("General job")


def written(day):
    return f"{day.day:02d} {_(MONTHS[day.month - 1])} {day.year}"


def range_text(data):
    start, end = (date.fromisoformat(data["range"][key]) for key in ("start", "end"))
    return _("{start} to {end}", start=written(start), end=written(end))


def local_date(iso, company):
    return datetime.fromisoformat(iso).replace(tzinfo=timezone.utc).astimezone(ZoneInfo(company.timezone)).date()


def blank(value):
    return "" if value is None else value


def amount(cents_value, currency):
    """Money as written in the files: "KSh 18,000" (no cents when there are none)."""
    if cents_value is None:
        return ""
    whole = cents_value / 100
    text = f"{whole:,.0f}" if cents_value % 100 == 0 else f"{whole:,.2f}"
    return f"KSh {text}" if currency == "KES" else f"{currency} {text}"


def summary_rows(data):
    jobs, hours, money = data["jobs"], data["hours"], data["money"]
    rows = [
        (_("Jobs created"), jobs["created"]),
        (_("Jobs completed"), jobs["completed"]),
        (_("Jobs cancelled"), jobs["cancelled"]),
        (_("Median time to complete (hours)"), blank(jobs["median_hours_to_complete"])),
        (_("Average time to complete (hours)"), blank(jobs["average_hours_to_complete"])),
        (_("Average hours worked per completed job"), blank(jobs["average_hours_worked"])),
        (_("Hours clocked in"), hours["clocked"]),
        (_("Hours on jobs"), hours["on_jobs"]),
        (_("Equipment with a repeat fault"), len(data["repeat_faults"])),
    ]
    # Money only on plans with job costing.
    if money:
        rows += [
            (_("Cost of completed jobs"), amount(money["cost_cents"], money["currency"])),
            (_("Charged for {count} of them", count=money["charged_jobs"]), amount(money["charged_cents"], money["currency"])),
            (_("Margin on the jobs charged"), amount(money["margin_cents"], money["currency"])),
        ]
    return rows


def technician_headings():
    return [_("Technician"), _("Hours clocked in"), _("Hours on jobs"), _("Share of time on jobs (%)"), _("Jobs completed"), _("Median time to complete (hours)"), _("Rating"), _("Reviews")]


def technician_row(row):
    return [
        row["name"],
        row["clocked_hours"],
        row["job_hours"],
        blank(row["share_on_jobs"]),
        row["jobs_completed"],
        blank(row["median_hours_to_complete"]),
        blank(row["rating"]),
        row["reviews"],
    ]


def to_xlsx(data, company):
    zone = ZoneInfo(company.timezone)
    book = Workbook()

    def sheet(title, headings, rows, first=False):
        page = book.active if first else book.create_sheet()
        page.title = clean_title(title)
        page.append(headings)
        for cell in page[1]:
            cell.font = Font(bold=True)
            cell.fill = HEAD_FILL
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in rows:
            page.append(row)
            # Text people typed (a job title "=HYPERLINK(...)") stays text: Excel never runs it.
            for cell in page[page.max_row]:
                if cell.data_type == "f":
                    cell.data_type = "s"
        page.freeze_panes = "A2"
        for index, heading in enumerate(headings, start=1):
            longest = max([len(str(heading))] + [len(str(row[index - 1])) for row in rows if index - 1 < len(row)])
            page.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 48)
        return page

    summary = sheet(_("Summary"), [_("Insights for {company}", company=company.name), range_text(data)], summary_rows(data), first=True)
    summary.column_dimensions["A"].width = 44

    sheet(_("Technicians"), technician_headings(), [technician_row(row) for row in data["technicians"]])

    jobs = [item["job"] for item in data["_completed_jobs"]]
    customers = {c.id: c.name for c in Customer.query.filter(Customer.id.in_([j.customer_id for j in jobs if j.customer_id] or [0])).all()}
    machines = {m.id: m.name for m in Machine.query.filter(Machine.id.in_([j.machine_id for j in jobs if j.machine_id] or [0])).all()}
    people = {u.id: u.name for u in User.query.filter(User.id.in_([j.mechanic_id for j in jobs if j.mechanic_id] or [0])).all()}

    def local(moment):
        return moment.replace(tzinfo=timezone.utc).astimezone(zone).replace(tzinfo=None)

    money = data["money"] is not None
    completed = sheet(
        _("Completed jobs"),
        [
            _("Job"), _("Title"), _("Customer"), _("Equipment"), _("Technician"), _("Job type"), _("Created"), _("Completed"),
            _("Hours to complete"), _("Hours worked"),
            *([_("Labour cost"), _("Parts cost"), _("Charged"), _("Margin")] if money else []),
        ],
        [
            [
                item["job"].id,
                item["job"].title,
                customers.get(item["job"].customer_id, ""),
                machines.get(item["job"].machine_id, ""),
                people.get(item["job"].mechanic_id, ""),
                item["job"].job_type or _(GENERAL),
                local(item["job"].created_at),
                local(item["completed_at"]),
                round(item["hours_to_complete"] or 0, 1),
                round(item["hours_worked"], 1),
                *[
                    item["costing"][key] / 100 if item["costing"][key] is not None else None
                    for key in (("labour_cents", "parts_cents", "charged_cents", "margin_cents") if money else ())
                ],
            ]
            for item in data["_completed_jobs"]
        ],
    )
    for column in ("G", "H"):
        completed.column_dimensions[column].width = 18
        for cell in completed[column][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm"
    # Money in the company's currency, as numbers Excel can add up.
    for column in ("K", "L", "M", "N") if money else ():
        for cell in completed[column][1:]:
            cell.number_format = "#,##0.00"

    sheet(
        _("Completed over time"),
        [_("Week starting") if data["range"]["unit"] == "week" else _("Day"), _("Jobs completed")],
        [[date.fromisoformat(point["start"]), point["count"]] for point in data["completed_over_time"]],
    )
    sheet(_("Job types"), [_("Job type"), _("Jobs completed")], [[row["job_type"] or _(GENERAL), row["count"]] for row in data["by_type"]])
    sheet(
        _("Repeat faults"),
        [_("Equipment"), _("Type"), _("Customer"), _("Faults in 90 days"), _("Last fault")],
        [
            [row["name"], row["equipment_type"] or "", row["customer"] or "", row["faults"], local(datetime.fromisoformat(row["last_fault"])) if row["last_fault"] else ""]
            for row in data["repeat_faults"]
        ],
    )
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def shown(value):
    """A table cell: 0 stays 0; nothing is a dash."""
    return "-" if value in ("", None) else clean(str(value))


def clean_title(title):
    """Sheet names: at most 31 characters, none of []:*?/\\."""
    return "".join(char for char in title if char not in "[]:*?/\\")[:31]


class InsightsPdf(FPDF):
    def __init__(self, company_name, subtitle):
        super().__init__(format="A4")
        self.company_name = company_name
        self.subtitle = subtitle
        self.set_auto_page_break(True, margin=18)
        self.set_margins(16, 16, 16)

    def header(self):
        self.set_font("Helvetica", "B", 13)
        self.set_text_color(*INK)
        self.cell(0, 7, clean(self.company_name), new_x="LMARGIN", new_y="NEXT")
        self.set_font("Helvetica", "", 9)
        self.set_text_color(*MUTED)
        self.cell(0, 5, clean(_("Insights") + f"  -  {self.subtitle}"), new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(*LINE)
        self.line(16, self.get_y() + 2, 194, self.get_y() + 2)
        self.ln(6)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(*MUTED)
        self.cell(0, 5, _("Page {number}", number=self.page_no()) + "/{nb}", align="R")

    def section(self, title, room=40):
        if self.get_y() > 280 - room:
            self.add_page()
        self.ln(3)
        self.set_font("Helvetica", "B", 11)
        self.set_text_color(*ACCENT)
        self.cell(0, 7, clean(title), new_x="LMARGIN", new_y="NEXT")
        self.set_text_color(*INK)

    def table(self, headings, rows, widths):
        self.set_font("Helvetica", "B", 8)
        self.set_text_color(*MUTED)
        for heading, width in zip(headings, widths):
            self.cell(width, 6, clean(heading), border="B")
        self.ln()
        self.set_font("Helvetica", "", 9)
        self.set_text_color(*INK)
        for row in rows:
            if self.get_y() > 270:
                self.add_page()
            for value, width in zip(row, widths):
                self.cell(width, 6, shown(value))
            self.ln()

    def columns(self, points, label):
        """A column chart: one series, so no legend; the largest value is labelled on its cap."""
        if not points:
            return
        left, top, width, height = 16, self.get_y() + 4, 178, 42
        most = max(point[1] for point in points) or 1
        slot = width / len(points)
        bar = min(slot * 0.6, 8)
        self.set_draw_color(*LINE)
        self.line(left, top + height, left + width, top + height)
        self.set_fill_color(*ACCENT)
        peak = max(range(len(points)), key=lambda index: points[index][1])
        for index, (name, value) in enumerate(points):
            x = left + index * slot + (slot - bar) / 2
            tall = height * value / most
            if tall:
                self.rect(x, top + height - tall, bar, tall, style="F")
            if index == peak and value:
                self.set_font("Helvetica", "B", 7)
                self.set_text_color(*INK)
                self.set_xy(x - 6, top + height - tall - 5)
                self.cell(bar + 12, 4, str(value), align="C")
        # First and last labels under the axis.
        self.set_font("Helvetica", "", 7)
        self.set_text_color(*MUTED)
        self.set_xy(left, top + height + 1)
        self.cell(width / 2, 4, clean(points[0][0]))
        self.set_xy(left + width / 2, top + height + 1)
        self.cell(width / 2, 4, clean(points[-1][0]), align="R")
        self.set_xy(left, top + height + 6)
        self.set_text_color(*INK)
        self.set_font("Helvetica", "", 8)
        self.cell(0, 4, clean(label), new_x="LMARGIN", new_y="NEXT")


def to_pdf(data, company):
    pdf = InsightsPdf(company.name, range_text(data))
    pdf.add_page()

    pdf.section(_("Summary"))
    pdf.set_font("Helvetica", "", 10)
    for label, value in summary_rows(data):
        pdf.set_text_color(*MUTED)
        pdf.cell(110, 6, clean(label))
        pdf.set_text_color(*INK)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 6, shown(value), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)

    pdf.section(_("Jobs completed"), room=70)
    weekly = data["range"]["unit"] == "week"
    points = [(written(date.fromisoformat(point["start"])), point["count"]) for point in data["completed_over_time"]]
    pdf.columns(points, _("A week a column, by the Monday it starts") if weekly else _("A day a column"))

    pdf.section(_("Technicians"))
    headings = technician_headings()
    pdf.table(
        [headings[0], _("Clocked"), _("On jobs"), _("On jobs %"), _("Completed"), _("Median h"), _("Rating")],
        [technician_row(row)[:7] for row in data["technicians"]],
        [50, 20, 20, 20, 22, 22, 18],
    )

    pdf.section(_("Job types"))
    pdf.table([_("Job type"), _("Jobs completed")], [[row["job_type"] or _(GENERAL), row["count"]] for row in data["by_type"]] or [[_("No jobs completed"), ""]], [120, 40])

    pdf.section(_("Repeat faults"))
    if data["repeat_faults"]:
        pdf.table(
            [_("Equipment"), _("Customer"), _("Faults in 90 days"), _("Last fault")],
            [
                [row["name"], row["customer"] or "", row["faults"], written(local_date(row["last_fault"], company)) if row["last_fault"] else ""]
                for row in data["repeat_faults"]
            ],
            [62, 56, 30, 30],
        )
    else:
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 6, clean(_("No equipment needed fixing twice in 90 days.")), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
